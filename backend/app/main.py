import copy
import csv
import hashlib
import hmac
import io
import json
import os
import secrets
import time
import uuid
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File, Form, Query
from fastapi.responses import FileResponse, StreamingResponse, JSONResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func, delete
from sqlalchemy.exc import IntegrityError
from sqlalchemy.exc import OperationalError
from .database import Base, engine, SessionLocal, Record, User, AuthSession, Audit, Document, Config, ImportBatch, now
from .catalog import CATALOG, ROLES, ADMIN, CONTRACTOR_ROLES, DEFAULT_SETTINGS
from .security import db_session, current_user, hash_password, verify_password, scope_query, can_write, require_admin, public_user
from .business import fail, get_record, serialize, save_record, transition, audit, settings, validate, natural_key
from .analytics import dashboard

STORAGE = Path(os.getenv("MCMS_STORAGE", "storage"))
DEMO = os.getenv("MCMS_DEMO", "false").lower() == "true"
attempts = defaultdict(deque)
DEFAULT_BRANDING = {
    "application_name": "MOne CoalChain",
    "application_tagline": "MINING OPERATIONS",
    "logo_version": 0,
}

def branding(db):
    entry = db.get(Config, "branding")
    value = {**DEFAULT_BRANDING, **(entry.value if entry else {})}
    if (STORAGE / "branding-logo").is_file() and value.get("logo_content_type"):
        value["logo_url"] = f"/api/branding/logo?v={value['logo_version']}"
    else:
        value["logo_url"] = None
    return value

@asynccontextmanager
async def lifespan(app):
    # MySQL's first-run entrypoint briefly exposes a temporary server while it
    # creates the schema, then restarts the real server. Compose can observe a
    # successful health probe during that window, so retry the initial schema
    # connection rather than crashing the API startup.
    for attempt in range(30):
        try:
            Base.metadata.create_all(engine)
            break
        except OperationalError:
            if attempt == 29:
                raise
            await __import__("asyncio").sleep(2)
    STORAGE.mkdir(parents=True, exist_ok=True)
    with SessionLocal() as db:
        if not db.scalar(select(User).limit(1)):
            password = os.getenv("MCMS_ADMIN_PASSWORD", "")
            if len(password) < 12:
                raise RuntimeError("MCMS_ADMIN_PASSWORD must contain at least 12 characters.")
            db.add(User(id="USR-ADMIN", email="admin@coalchain.local", name="Administrator", role=ADMIN, password_hash=hash_password(password)))
            db.commit()
        if not db.get(Config, "settings"):
            db.add(Config(key="settings", value=copy.deepcopy(DEFAULT_SETTINGS)))
            db.commit()
        if not db.get(Config, "branding"):
            db.add(Config(key="branding", value=copy.deepcopy(DEFAULT_BRANDING)))
            db.commit()
        if DEMO and not db.scalar(select(Record).limit(1)):
            from .seed import seed
            seed(db)
    yield

app = FastAPI(title="MOne CoalChain API", version="0.1.0", lifespan=lifespan, docs_url="/api/docs", openapi_url="/api/openapi.json", redoc_url=None)

@app.middleware("http")
async def protections(request, call_next):
    if request.method in ("POST", "PUT", "PATCH", "DELETE") and not request.url.path.startswith("/api/internal/"):
        # A custom header forces a CORS preflight; no foreign origins are granted CORS.
        if request.headers.get("x-mcms-request") != "1":
            return JSONResponse({"detail": "Header keamanan tidak valid."}, status_code=403)
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["Cache-Control"] = "no-store"
    return response

@app.exception_handler(IntegrityError)
async def integrity_error(request, exc):
    return JSONResponse({"detail": "Konflik data atau transaksi duplikat. Muat ulang dan periksa referensi."}, status_code=409)

class Login(BaseModel):
    email: str = Field(max_length=150)
    password: str = Field(max_length=200)

class Mutation(BaseModel):
    data: dict
    request_id: str | None = Field(default=None, max_length=100)
    version: int | None = None
    reason: str = Field(default="", max_length=10000)

class Action(BaseModel):
    action: str
    comment: str = Field(default="", max_length=10000)
    version: int

@app.get("/api/health")
def health():
    return {"status": "ok", "service": "mcms-api"}

@app.get("/api/branding")
def get_branding(db=Depends(db_session)):
    return branding(db)

@app.get("/api/branding/logo")
def get_branding_logo(db=Depends(db_session)):
    value = branding(db)
    path = STORAGE / "branding-logo"
    if not value["logo_url"] or not path.is_file():
        fail("Logo belum dikonfigurasi.", 404)
    return FileResponse(path, media_type=value["logo_content_type"], filename="application-logo")

@app.post("/api/auth/login")
def login(body: Login, request: Request, response: Response, db=Depends(db_session)):
    key = ((request.client.host if request.client else "unknown"), body.email.lower())
    bucket = attempts[key]
    cutoff = time.monotonic() - 900
    while bucket and bucket[0] < cutoff:
        bucket.popleft()
    if len(bucket) >= 10:
        fail("Terlalu banyak percobaan. Tunggu 15 menit.", 429)
    bucket.append(time.monotonic())
    user = db.scalar(select(User).where(User.email == body.email.lower()))
    if not user or not user.active or not verify_password(body.password, user.password_hash):
        fail("Email atau password tidak benar.", 401)
    bucket.clear()
    token = secrets.token_urlsafe(48)
    db.execute(delete(AuthSession).where(AuthSession.expires_at < now()))
    db.add(AuthSession(token_hash=hashlib.sha256(token.encode()).hexdigest(), user_id=user.id, expires_at=now()+timedelta(minutes=30)))
    audit(db, user, "LOGIN")
    db.commit()
    response.set_cookie("mcms_session", token, httponly=True, secure=os.getenv("MCMS_SECURE_COOKIE", "false") == "true", samesite="strict", path="/")
    return public_user(user)

@app.post("/api/auth/logout")
def logout(request: Request, response: Response, user=Depends(current_user), db=Depends(db_session)):
    db.execute(delete(AuthSession).where(AuthSession.token_hash == hashlib.sha256(request.cookies.get("mcms_session", "").encode()).hexdigest()))
    audit(db, user, "LOGOUT")
    db.commit()
    response.delete_cookie("mcms_session", path="/")
    return {"ok": True}

@app.get("/api/auth/me")
def me(user=Depends(current_user)):
    return public_user(user)

@app.post("/api/auth/password")
def password(body: dict, user=Depends(current_user), db=Depends(db_session)):
    new = str(body.get("new_password", ""))
    if not verify_password(str(body.get("current_password", "")), user.password_hash) or not 12 <= len(new) <= 200:
        fail("Password lama salah atau password baru kurang dari 12 karakter.")
    target = db.get(User, user.id)
    target.password_hash = hash_password(new)
    db.execute(delete(AuthSession).where(AuthSession.user_id == user.id))
    audit(db, user, "PASSWORD_CHANGED")
    db.commit()
    return {"ok": True}

@app.get("/api/catalog")
def catalog(user=Depends(current_user), db=Depends(db_session)):
    cfg = settings(db)
    return {"modules": {k: {**v, "can_write": can_write(user, k), "workflow": cfg["workflows"].get(k, [])} for k, v in CATALOG.items()}, "roles": ROLES, "demo": DEMO}

@app.get("/api/dashboard")
def dashboard_route(period: str | None=None, site: str | None=None, contractor: str | None=None, user=Depends(current_user), db=Depends(db_session)):
    try:
        if period:
            date.fromisoformat(period + "-01")
    except ValueError:
        fail("Periode tidak valid.")
    return dashboard(db, user, period, site, contractor)

@app.get("/api/lookups")
def lookups(user=Depends(current_user), db=Depends(db_session)):
    q = scope_query(user, select(Record).where(Record.status.not_in(["CANCELLED", "REJECTED"])), Record)
    return [{"id": r.id, "kind": r.kind, "name": r.data.get("name", r.id), "status": r.status,
             "site_id": r.site_id, "contractor_id": r.contractor_id, "date": r.data.get("date"), "activity": r.data.get("activity")} for r in db.scalars(q)]

def check_kind(kind):
    if kind not in CATALOG:
        fail("Modul tidak ditemukan.", 404)

@app.get("/api/records/{kind}")
def records(kind: str, search: str="", status: str="", site: str="", contractor: str="", period: str="",
            page: int=Query(1, ge=1), limit: int=Query(25, ge=1, le=100),
            user=Depends(current_user), db=Depends(db_session)):
    check_kind(kind)
    q = scope_query(user, select(Record).where(Record.kind == kind), Record)
    if status:
        q = q.where(Record.status == status)
    if site:
        q = q.where(Record.site_id == site)
    if contractor:
        q = q.where(Record.contractor_id == contractor)
    # Search across JSON values in a database-portable way for the initial release.
    items = [serialize(r) for r in db.scalars(q.order_by(Record.created_at.desc()))]
    if search:
        items = [r for r in items if search.lower() in json.dumps(r, ensure_ascii=False).lower()]
    if period:
        items = [r for r in items if str(r.get("date", r.get("period", r.get("start", "")))).startswith(period)]
    return {"items": items[(page-1)*limit:page*limit], "total": len(items), "page": page, "limit": limit}

@app.post("/api/records/{kind}", status_code=201)
def create_record(kind: str, body: Mutation, user=Depends(current_user), db=Depends(db_session)):
    check_kind(kind)
    request_key = "request:" + hashlib.sha256((user.id + ":" + kind + ":" + (body.request_id or "")).encode()).hexdigest()
    payload_hash = hashlib.sha256(json.dumps(body.data, sort_keys=True).encode()).hexdigest()
    existing = db.get(Config, request_key) if body.request_id else None
    if existing:
        if existing.value["payload_hash"] != payload_hash:
            fail("Request ID sudah dipakai untuk payload berbeda.", 409)
        return serialize(get_record(db, user, existing.value["record_id"]))
    row = save_record(db, user, kind, body.data)
    if body.request_id:
        db.add(Config(key=request_key, value={"record_id": row.id, "payload_hash": payload_hash}))
    db.commit()
    return serialize(row)

@app.get("/api/record/{record_id}")
def record_detail(record_id: str, user=Depends(current_user), db=Depends(db_session)):
    row = get_record(db, user, record_id)
    logs = list(db.scalars(select(Audit).where(Audit.record_id == row.id).order_by(Audit.id.desc())))
    docs = list(db.scalars(select(Document).where(Document.record_id == row.id)))
    return {"record": serialize(row), "audit": [audit_dict(a) for a in logs], "documents": [document_dict(d) for d in docs]}

@app.put("/api/record/{record_id}")
def update_record(record_id: str, body: Mutation, user=Depends(current_user), db=Depends(db_session)):
    row = get_record(db, user, record_id, lock=True)
    if body.version != row.version:
        fail("Data telah berubah. Muat ulang.", 409)
    result = save_record(db, user, row.kind, body.data, row, body.reason)
    db.commit()
    return serialize(result)

@app.post("/api/record/{record_id}/action")
def action(record_id: str, body: Action, user=Depends(current_user), db=Depends(db_session)):
    row = get_record(db, user, record_id, lock=True)
    transition(db, user, row, body.action, body.comment, body.version)
    db.commit()
    return serialize(row)

def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

@app.get("/api/export/{kind}")
def export(kind: str, user=Depends(current_user), db=Depends(db_session)):
    check_kind(kind)
    rows = list(db.scalars(scope_query(user, select(Record).where(Record.kind == kind), Record)))
    fields = ["id"] + [f["key"] for f in CATALOG[kind]["fields"]] + ["status", "created_by", "created_at"]
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fields, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe_cell(v) for k, v in serialize(row).items()})
    return Response("\ufeff" + output.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{kind}.csv"'})

@app.get("/api/import/{kind}/template")
def template(kind: str, user=Depends(current_user)):
    check_kind(kind)
    return Response(",".join(f["key"] for f in CATALOG[kind]["fields"]) + "\n", media_type="text/csv",
                    headers={"Content-Disposition": f'attachment; filename="{kind}-template.csv"'})

async def read_upload(upload, max_size=10*1024*1024):
    content = await upload.read(max_size + 1)
    if len(content) > max_size:
        fail("Ukuran file maksimum 10 MB.")
    return content

@app.post("/api/import/{kind}/preview")
async def import_preview(kind: str, file: UploadFile=File(...), user=Depends(current_user), db=Depends(db_session)):
    check_kind(kind)
    if not can_write(user, kind):
        fail("Tidak memiliki hak import.", 403)
    content = await read_upload(file)
    try:
        if (file.filename or "").lower().endswith(".xlsx"):
            import zipfile
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if sum(i.file_size for i in z.infolist()) > 30*1024*1024:
                    fail("Workbook terlalu besar setelah ekstraksi.")
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = wb.active
            it = sheet.iter_rows(values_only=True)
            headers = [str(v) for v in next(it)]
            raw = []
            for index, values in enumerate(it):
                if index >= 500:
                    fail("Maksimum 500 baris per import.")
                if not any(v is not None for v in values):
                    continue
                if any(isinstance(v, str) and v.startswith("=") for v in values):
                    fail("Formula Excel tidak diterima; gunakan nilai.")
                raw.append({k: (v.date().isoformat() if isinstance(v, datetime) else v) for k, v in zip(headers, values)})
            wb.close()
        elif (file.filename or "").lower().endswith(".csv"):
            raw = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            fail("Gunakan file CSV UTF-8 atau XLSX.")
    except HTTPException:
        raise
    except Exception:
        fail("File tidak dapat dibaca. Periksa format dan header template.")
    if not raw or len(raw) > 500:
        fail("Import harus berisi 1 sampai 500 baris.")
    errors, keys = [], set()
    for index, payload in enumerate(raw):
        try:
            data = validate(db, user, kind, payload)
            key = natural_key(kind, data)
            if key and (key in keys or db.scalar(select(Record).where(Record.natural_key == key))):
                fail("Transaksi duplikat.")
            if key:
                keys.add(key)
        except HTTPException as exc:
            errors.append({"row": index + 2, "message": exc.detail})
    batch = ImportBatch(id=uuid.uuid4().hex, user_id=user.id, kind=kind, rows=raw)
    db.add(batch)
    db.commit()
    return {"batch_id": batch.id, "rows": raw[:50], "count": len(raw), "errors": errors}

@app.post("/api/import/confirm/{batch_id}")
def import_confirm(batch_id: str, user=Depends(current_user), db=Depends(db_session)):
    batch = db.scalar(select(ImportBatch).where(ImportBatch.id == batch_id).with_for_update())
    if not batch or batch.user_id != user.id:
        fail("Batch tidak ditemukan.", 404)
    if batch.posted:
        fail("Batch sudah diposting.", 409)
    if batch.created_at < now() - timedelta(hours=2):
        fail("Preview kedaluwarsa. Upload kembali.", 409)
    # Revalidate all rows in one transaction. Any failure rolls back the whole batch.
    for payload in batch.rows:
        save_record(db, user, batch.kind, payload)
    batch.posted = 1
    db.commit()
    return {"posted": len(batch.rows)}

def document_dict(d):
    return {"id": d.id, "name": d.name, "version": d.version, "document_type": d.document_type,
            "uploaded_by": d.uploaded_by, "uploaded_at": d.uploaded_at.isoformat()}

@app.post("/api/record/{record_id}/documents")
async def upload_document(record_id: str, file: UploadFile=File(...), document_type: str=Form("Supporting Evidence"),
                          user=Depends(current_user), db=Depends(db_session)):
    row = get_record(db, user, record_id, lock=True)
    if not can_write(user, row.kind) or row.status not in ("DRAFT", "RETURNED", "REJECTED", "ACTIVE"):
        fail("Lampiran hanya dapat ditambahkan oleh editor sebelum submit.", 403)
    name = Path((file.filename or "document").replace("\\", "/")).name[:200]
    ext = Path(name).suffix.lower()
    if ext not in (".pdf", ".png", ".jpg", ".jpeg", ".csv", ".xlsx", ".docx", ".txt"):
        fail("Tipe file tidak didukung.")
    content = await read_upload(file)
    version = (db.scalar(select(func.max(Document.version)).where(Document.record_id == row.id, Document.name == name)) or 0) + 1
    doc = Document(id=uuid.uuid4().hex, record_id=row.id, name=name, version=version, document_type=document_type[:100],
                   content_type="application/octet-stream", uploaded_by=user.email)
    path = STORAGE / doc.id
    path.write_bytes(content)
    try:
        db.add(doc)
        audit(db, user, "DOCUMENT_UPLOAD", row, reason=name, after={"document_id": doc.id, "version": version})
        db.commit()
    except Exception:
        path.unlink(missing_ok=True)
        raise
    return document_dict(doc)

@app.get("/api/documents/{document_id}")
def download_document(document_id: str, user=Depends(current_user), db=Depends(db_session)):
    doc = db.get(Document, document_id)
    if not doc:
        fail("Dokumen tidak ditemukan.", 404)
    get_record(db, user, doc.record_id)
    path = STORAGE / doc.id
    if not path.is_file():
        fail("File dokumen tidak tersedia.", 404)
    return FileResponse(path, filename=doc.name, media_type=doc.content_type)

def audit_dict(a):
    return {"id": a.id, "record_id": a.record_id, "actor": a.actor, "action": a.action, "reason": a.reason,
            "before": a.before, "after": a.after, "timestamp": a.timestamp.isoformat()}

@app.get("/api/audit")
def audit_logs(page: int=Query(1, ge=1), user=Depends(current_user), db=Depends(db_session)):
    allowed = scope_query(user, select(Record.id), Record)
    q = select(Audit).where((Audit.record_id.in_(allowed)) | ((Audit.record_id == None) & (Audit.actor == user.email)))
    if user.role == ADMIN:
        q = select(Audit)
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    items = list(db.scalars(q.order_by(Audit.id.desc()).offset((page-1)*50).limit(50)))
    return {"items": [audit_dict(a) for a in items], "total": total}

@app.get("/api/settings")
def get_settings(user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    return settings(db)

@app.put("/api/settings")
def put_settings(body: dict, user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    required = set(DEFAULT_SETTINGS)
    if set(body) != required:
        fail("Konfigurasi tidak lengkap.")
    for key in required - {"weights", "workflows"}:
        if not isinstance(body[key], (float, int)) or not 0 <= body[key] <= 10000:
            fail(f"Nilai {key} tidak valid.")
    weights = body["weights"]
    if set(weights) != set(DEFAULT_SETTINGS["weights"]) or any(not isinstance(v, (int, float)) or not 0 <= v <= 100 for v in weights.values()) or abs(sum(weights.values())-100) > .001:
        fail("Total bobot KPI harus 100%.")
    if set(body["workflows"]) != set(DEFAULT_SETTINGS["workflows"]):
        fail("Daftar workflow tidak lengkap.")
    for kind, roles in body["workflows"].items():
        if not isinstance(roles, list) or not 1 <= len(roles) <= 10 or any(r not in ROLES for r in roles):
            fail(f"Workflow {kind} tidak valid.")
    config = db.get(Config, "settings")
    before = config.value
    config.value = body
    audit(db, user, "SETTINGS_CHANGED", before=before, after=body)
    db.commit()
    return body

@app.put("/api/settings/branding")
def put_branding(body: dict, user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    application_name = str(body.get("application_name", "")).strip()
    application_tagline = str(body.get("application_tagline", "")).strip().upper()
    if not 1 <= len(application_name) <= 60 or not 1 <= len(application_tagline) <= 60:
        fail("Nama aplikasi dan label operasional harus berisi 1 sampai 60 karakter.")
    if any(ord(char) < 32 for char in application_name + application_tagline):
        fail("Nama aplikasi atau label operasional tidak valid.")
    entry = db.get(Config, "branding")
    before = branding(db)
    entry.value = {**entry.value, "application_name": application_name, "application_tagline": application_tagline}
    audit(db, user, "BRANDING_UPDATED", reason="Branding configuration", before=before, after=branding(db))
    db.commit()
    return branding(db)

def image_type(content: bytes):
    if content.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if content.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if len(content) >= 12 and content[:4] == b"RIFF" and content[8:12] == b"WEBP":
        return "image/webp"
    return None

@app.post("/api/settings/branding/logo")
async def upload_branding_logo(file: UploadFile=File(...), user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    content = await read_upload(file, max_size=2 * 1024 * 1024)
    content_type = image_type(content)
    if not content_type:
        fail("Logo harus berupa gambar PNG, JPG/JPEG, atau WebP yang valid.")
    entry = db.get(Config, "branding")
    before = branding(db)
    temporary = STORAGE / "branding-logo-upload"
    temporary.write_bytes(content)
    temporary.replace(STORAGE / "branding-logo")
    entry.value = {**entry.value, "logo_content_type": content_type, "logo_version": int(entry.value.get("logo_version", 0)) + 1}
    audit(db, user, "BRANDING_LOGO_UPLOADED", reason=Path(file.filename or "application-logo").name, before=before, after=branding(db))
    db.commit()
    return branding(db)

@app.delete("/api/settings/branding/logo")
def remove_branding_logo(user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    entry = db.get(Config, "branding")
    before = branding(db)
    (STORAGE / "branding-logo").unlink(missing_ok=True)
    entry.value = {key: value for key, value in entry.value.items() if key != "logo_content_type"}
    entry.value["logo_version"] = int(entry.value.get("logo_version", 0)) + 1
    audit(db, user, "BRANDING_LOGO_REMOVED", reason="Logo removed", before=before, after=branding(db))
    db.commit()
    return branding(db)

@app.get("/api/users")
def users(user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    return [public_user(u) for u in db.scalars(select(User))]

@app.post("/api/users")
def create_user(body: dict, user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    email, name, role, password = str(body.get("email", "")).lower().strip(), str(body.get("name", "")).strip(), body.get("role"), str(body.get("password", ""))
    if "@" not in email or len(email) > 150 or not 1 <= len(name) <= 150 or role not in ROLES or not 12 <= len(password) <= 200:
        fail("Nama, email, role, atau password tidak valid (minimum 12 karakter).")
    site_id, contractor_id = body.get("site_id") or None, body.get("contractor_id") or None
    if site_id:
        get_record(db, user, site_id, "sites")
    if contractor_id:
        contractor = get_record(db, user, contractor_id, "contractors")
        if contractor.site_id != site_id:
            fail("Site akun harus sesuai kontraktor.")
    if role in CONTRACTOR_ROLES and (not contractor_id or not site_id):
        fail("Akun kontraktor wajib memiliki scope site dan contractor.")
    new = User(id="USR-"+uuid.uuid4().hex[:12], email=email, name=name, role=role, password_hash=hash_password(password), site_id=site_id, contractor_id=contractor_id)
    db.add(new)
    audit(db, user, "USER_CREATED", after=public_user(new))
    db.commit()
    return public_user(new)

@app.post("/api/users/{user_id}/toggle")
def toggle_user(user_id: str, user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    target = db.get(User, user_id)
    if not target or target.id == user.id:
        fail("Tidak dapat mengubah status akun ini.")
    target.active = 0 if target.active else 1
    db.execute(delete(AuthSession).where(AuthSession.user_id == target.id))
    audit(db, user, "USER_STATUS_CHANGED", after=public_user(target))
    db.commit()
    return public_user(target)

@app.post("/api/internal/refresh-alerts")
def refresh_alerts(request: Request, db=Depends(db_session)):
    expected = os.getenv("MCMS_WORKER_SECRET", "")
    if not expected or not hmac.compare_digest(request.headers.get("authorization", ""), "Bearer " + expected):
        fail("Worker tidak terautentikasi.", 401)
    admin = db.scalar(select(User).where(User.role == ADMIN, User.active == 1))
    snapshot = {"refreshed_at": now().isoformat(), "count": len(dashboard(db, admin)["alerts"])}
    entry = db.get(Config, "worker_status")
    if entry:
        entry.value = snapshot
    else:
        db.add(Config(key="worker_status", value=snapshot))
    db.commit()
    return snapshot

@app.get("/api/integrations")
def integrations(user=Depends(current_user), db=Depends(db_session)):
    require_admin(user)
    entry = db.get(Config, "worker_status")
    return {"worker": entry.value if entry else None, "api_docs": "/api/docs",
            "connectors": [{"name": name, "status": "Belum dikonfigurasi"} for name in ["ERP / Odoo / SAP", "Weighbridge", "GPS / Fleet", "Fuel station", "Laboratory", "HR / Procurement"]],
            "authentication": "Session cookie; X-MCMS-Request: 1 untuk mutasi. Service account API key adalah pengembangan lanjutan."}

